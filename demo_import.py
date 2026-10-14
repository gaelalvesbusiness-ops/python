"""Offline end-to-end demonstration of the verified-number import (Workflow B).

Runs the real ``main.py`` CLI as a subprocess against a temporary SQLite database, using the
synthetic files in ``examples/import``. The temporary database first receives one simulated
website page (in memory, no network) so that preservation of website records and page progress
can be checked. Nothing contacts HonestPool or its API, and the real database is never opened.

    python demo_import.py            # prints PASS/FAIL per check; exit code 1 if any check fails
"""

from __future__ import annotations

import csv
import hashlib
import shutil
import sqlite3
import subprocess
import sys
import tempfile
from collections.abc import Sequence
from pathlib import Path

from openpyxl import load_workbook

from scraper.config import DATA_DIR, DB_PATH, PROJECT_ROOT, RANGE_START, SEGMENT_SIZE
from scraper.database import Database
from scraper.runner import plan_run, process_page
from tests.html_fixtures import FakeSite

SOURCE = "synthetic demonstration data"
EXAMPLES = PROJECT_ROOT / "examples" / "import"
TXT_NAME, CSV_NAME = "synthetic_verified_numbers.txt", "synthetic_verified_numbers.csv"
WEBSITE_NUMBERS = [1, 2]
COLUMNS = ["Number", "Start hex", "End hex", "Status"]
VERIFIED = "Проверено"
PROTECTED = [DB_PATH, DATA_DIR / "verified_records.csv", DATA_DIR / "verified_records.xlsx"]


class Demo:
    def __init__(self) -> None:
        self.failures = 0

    def check(self, label: str, ok: bool, detail: str = "") -> None:
        self.failures += not ok
        print(f"  [{'PASS' if ok else 'FAIL'}] {label}" + (f" - {detail}" if detail else ""))

    def cli(self, *args: str, label: str) -> subprocess.CompletedProcess[str]:
        command = [sys.executable, "main.py", *args]
        shown = " ".join(a if " " not in a else repr(a) for a in ["python", "main.py", *args])
        print(f"\n$ {shown}    # {label}")
        proc = subprocess.run(command, cwd=PROJECT_ROOT, capture_output=True, text=True, timeout=120,
                              stdin=subprocess.DEVNULL)
        for line in proc.stdout.splitlines():
            print(f"    {line}")
        for line in proc.stderr.splitlines():
            if " INFO " not in line:
                print(f"    {line}")
        print(f"    (exit code {proc.returncode})")
        return proc


def formula(number: int) -> tuple[str, str]:
    start = RANGE_START + (number - 1) * SEGMENT_SIZE
    return f"{start:018x}", f"{start + SEGMENT_SIZE:018x}"


def numbers_in(path: Path) -> list[int]:
    if path.suffix == ".txt":
        return [int(line) for line in path.read_text(encoding="utf-8").split() if line.strip()]
    with path.open(encoding="utf-8-sig", newline="") as fh:
        return [int(row["Number"]) for row in csv.DictReader(fh)]


def table(db_path: Path, sql: str) -> list[tuple]:
    conn = sqlite3.connect(f"{db_path.resolve().as_uri()}?mode=ro", uri=True)
    try:
        return [tuple(r) for r in conn.execute(sql)]
    finally:
        conn.close()


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def protected_state() -> dict[str, str | None]:
    return {str(p): sha256(p) if p.exists() else None for p in PROTECTED}


def read_csv(path: Path) -> list[list[str]]:
    with path.open(encoding="utf-8-sig", newline="") as fh:
        return list(csv.reader(fh))


def check_rows(demo: Demo, label: str, header: list, rows: list[list], expected: list[int]) -> None:
    demo.check(f"{label}: the four required columns", header == COLUMNS, f"{header}")
    numbers = [int(r[0]) for r in rows]
    demo.check(f"{label}: no duplicated segment numbers", len(numbers) == len(set(numbers)),
               f"{len(numbers)} rows")
    demo.check(f"{label}: expected records", numbers == expected, f"{numbers}")
    bad = [r for r in rows if (str(r[1]), str(r[2])) != formula(int(r[0])) or len(str(r[1])) != 18
           or len(str(r[2])) != 18 or r[3] != VERIFIED]
    demo.check(f"{label}: complete 18-digit hex matching the formula, Status {VERIFIED}", not bad,
               f"wrong rows {bad[:2]}" if bad else f"{len(rows)} rows")


def run(work: Path) -> int:
    demo = Demo()
    db_path = work / "demo.db"
    assert db_path.resolve() != DB_PATH.resolve()
    before = protected_state()
    txt, csv_file = work / TXT_NAME, work / CSV_NAME
    shutil.copyfile(EXAMPLES / TXT_NAME, txt)
    shutil.copyfile(EXAMPLES / CSV_NAME, csv_file)
    txt_numbers, csv_numbers = numbers_in(txt), numbers_in(csv_file)

    print(f"Temporary database: {db_path}")
    print(f"Simulated website page 1 with verified numbers {WEBSITE_NUMBERS} (in memory, no network); "
          f"pages 2-3 planned and pending")
    with Database(db_path) as db:
        plan_run(db, 1, 3)
        process_page(db, 1, FakeSite(verified={1: WEBSITE_NUMBERS}))
    pages_sql = "SELECT * FROM page_progress ORDER BY page_number"
    website_sql = "SELECT * FROM verified_records ORDER BY number"
    pages_before, website_before = table(db_path, pages_sql), table(db_path, website_sql)
    run_state_before = table(db_path, "SELECT * FROM run_state ORDER BY key")

    print("\n== PART B: offline import through main.py ==")
    proc = demo.cli("--db", str(db_path), "--import-verified", str(txt), "--source", SOURCE, label="TXT import")
    txt_unique = sorted(set(txt_numbers))
    demo.check("TXT import accepted", proc.returncode == 0)
    demo.check("TXT duplicates removed before storing",
               f"{len(txt_numbers)} valid, {len(txt_numbers) - len(txt_unique)} duplicate(s) removed, "
               f"{len(txt_unique)} unique" in proc.stdout)

    proc = demo.cli("--db", str(db_path), "--import-verified", str(csv_file), "--source", SOURCE, label="CSV import")
    csv_unique = sorted(set(csv_numbers))
    demo.check("CSV import accepted", proc.returncode == 0)
    demo.check("CSV: number already stored from the website is skipped, overlap with the TXT is not imported twice",
               "skipped 1 already stored from the website and 1 already imported" in proc.stdout)

    imported = sorted((set(txt_unique) | set(csv_unique)) - set(WEBSITE_NUMBERS))
    rows = table(db_path, "SELECT number, start_hex, end_hex, status FROM imported_records ORDER BY number")
    demo.check("imported numbers stored once each", [r[0] for r in rows] == imported, f"{[r[0] for r in rows]}")
    demo.check("hex values equal 0x400000000000000000 + (N - 1) * 2**44 and start + 2**44",
               all((r[1], r[2]) == formula(r[0]) for r in rows),
               f"e.g. {rows[0][0]}: {rows[0][1]}..{rows[0][2]}; {rows[-1][0]}: {rows[-1][1]}..{rows[-1][2]}")
    demo.check(f"status is {VERIFIED}", {r[3] for r in rows} == {VERIFIED})

    proc = demo.cli("--db", str(db_path), "--import-verified", str(txt), "--source", SOURCE,
                    label="same TXT again")
    demo.check("re-import adds nothing", proc.returncode == 0 and "Imported 0 new verified record(s)" in proc.stdout
               and f"{len(txt_unique)} already imported" in proc.stdout)

    counts_sql = ("SELECT (SELECT COUNT(*) FROM imported_records), (SELECT COUNT(*) FROM import_batches), "
                  "(SELECT COUNT(*) FROM verified_records)")
    counts = table(db_path, counts_sql)
    invalid_txt = work / "invalid_numbers.txt"
    invalid_txt.write_text("3001\n3002\n12a\n67108865\n", encoding="utf-8")
    proc = demo.cli("--db", str(db_path), "--import-verified", str(invalid_txt), "--source", SOURCE,
                    label="TXT with invalid numbers")
    demo.check("invalid numbers reject the whole file (exit code 4, line numbers reported)",
               proc.returncode == 4 and "line 3" in proc.stderr and "line 4" in proc.stderr)
    invalid_csv = work / "unverified.csv"
    invalid_csv.write_text(f"Number,Status\n3001,{VERIFIED}\n3002,Не проверено\n", encoding="utf-8")
    proc = demo.cli("--db", str(db_path), "--import-verified", str(invalid_csv), "--source", SOURCE,
                    label="CSV with an unverified row")
    demo.check("a row that is not 'Проверено' rejects the whole file", proc.returncode == 4)
    proc = demo.cli("--db", str(db_path), "--import-verified", str(txt), label="no source, non-interactive")
    demo.check("import without a stated authorized source is refused", proc.returncode == 4)
    demo.check("rejected imports wrote nothing (valid 3001/3002 in the rejected files were not stored)",
               table(db_path, counts_sql) == counts
               and not table(db_path, "SELECT 1 FROM imported_records WHERE number IN (3001, 3002)"))

    batches = table(db_path, "SELECT id, source, file_name, file_sha256, file_format, numbers_in_file, "
                             "duplicates, inserted, already_scraped, already_imported FROM import_batches")
    print("\n  import_batches:")
    for b in batches:
        print(f"    batch {b[0]}: source={b[1]!r} file={b[2]} format={b[4]} sha256={b[3][:16]}... "
              f"unique={b[5]} duplicates={b[6]} inserted={b[7]} already_scraped={b[8]} already_imported={b[9]}")
    demo.check("provenance: source recorded for every batch", {b[1] for b in batches} == {SOURCE})
    demo.check("provenance: file name and SHA-256 match the imported files",
               [(b[2], b[3]) for b in batches] == [(TXT_NAME, sha256(txt)), (CSV_NAME, sha256(csv_file)),
                                                   (TXT_NAME, sha256(txt))])
    links = table(db_path, "SELECT DISTINCT import_id FROM imported_records ORDER BY import_id")
    demo.check("provenance: each imported record points to the batch that added it", links == [(1,), (2,)])
    demo.check("website records unchanged and kept apart from imports", table(db_path, website_sql) == website_before)
    demo.check("website page progress unchanged (no page marked completed by the import)",
               table(db_path, pages_sql) == pages_before, f"{[(p[0], p[1]) for p in pages_before]}")
    demo.check("saved run range unchanged", table(db_path, "SELECT * FROM run_state ORDER BY key") == run_state_before)

    print("\n== PART C: export validation ==")
    expected_all = sorted(set(imported) | set(WEBSITE_NUMBERS))
    exports = {
        "imported CSV": ("csv", "imported", imported),
        "imported XLSX": ("xlsx", "imported", imported),
        "all CSV": ("csv", "all", expected_all),
        "all XLSX": ("xlsx", "all", expected_all),
        "website CSV": ("csv", "website", WEBSITE_NUMBERS),
    }
    for label, (fmt, source, expected) in exports.items():
        out = work / f"export_{source}.{fmt}"
        proc = demo.cli("--db", str(db_path), "--export", fmt, "--export-source", source, "--output", str(out),
                        label=label)
        demo.check(f"{label}: exported", proc.returncode == 0 and f"Exported {len(expected)} verified" in proc.stdout)
        if fmt == "csv":
            data = read_csv(out)
            check_rows(demo, label, data[0], data[1:], expected)
        else:
            ws = load_workbook(out)["Verified"]
            values = [[c.value for c in row] for row in ws.iter_rows()]
            check_rows(demo, label, values[0], values[1:], expected)
            hex_cells = [c for row in ws.iter_rows(min_row=2) for c in row[1:3]]
            demo.check(f"{label}: hex cells stored as text (type 's', format '@')",
                       all(c.data_type == "s" and c.number_format == "@" and isinstance(c.value, str)
                           for c in hex_cells), f"{len(hex_cells)} cells")

    demo.check("real database and exports untouched", protected_state() == before)
    print(f"\nResult: {'all checks passed' if not demo.failures else f'{demo.failures} check(s) FAILED'}")
    return 1 if demo.failures else 0


def main(argv: Sequence[str] | None = None) -> int:
    with tempfile.TemporaryDirectory(prefix="honestpool-import-demo-") as tmp:
        return run(Path(tmp))


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
