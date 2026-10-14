"""Hands-on offline walkthrough of the client requirements (Step 12).

Runs the real ``main.py`` commands against a temporary database. The browser is the
test suite's in-memory Playwright stand-in serving synthetic listing pages, so the real
fetcher, pacing, parser, SQLite persistence and progress report are exercised without
any network access. The real database is never opened for writing.
"""

from __future__ import annotations

import csv
import hashlib
import logging
import random
import sqlite3
import sys
import tempfile
import zipfile
from pathlib import Path
from urllib.parse import parse_qs, urlsplit

from openpyxl import load_workbook

import main as cli
from scraper import browser
from scraper.browser import (
    AccessRestrictedError,
    BrowserSettings,
    CaptchaRequiredError,
    ListingFetcher,
    RateLimitedError,
    TransientFetchError,
)
from scraper.config import DATA_DIR, DB_PATH, PROJECT_ROOT
from scraper.database import Database
from scraper.pagination import page_segments
from scraper.parser import segment_range
from scraper.progress import ProgressReporter
from scraper.runner import plan_run, run_pages
from tests.fake_playwright import (
    BLANK_URL,
    FakePage,
    FakePlaywright,
    FakeResponse,
    captcha,
    factory,
    http,
    success,
    timeout,
)
from tests.html_fixtures import CAPTCHA_PAGE, listing_html, listing_page, listing_row

BASE = "https://honestpool.ru/puzzle?id=71"
VERIFIED = {
    1000: [99_901, 99_950], 1001: [], 1002: [100_150], 1003: [100_201, 100_300], 1004: [], 1005: [100_401],
    2000: [199_999], 2001: [200_050], 2002: [], 2003: [200_201],
}
PACING = ["--delay", "5", "--jitter", "0"]
PROTECTED = [DB_PATH, DATA_DIR / "verified_records.csv", DATA_DIR / "verified_records.xlsx",
             PROJECT_ROOT / "delivery" / "honestpool-scraper-demo-20261013.zip"]

results: list[tuple[str, str, str]] = []


# --- output helpers ---------------------------------------------------------------------------

def section(title: str) -> None:
    print(f"\n{'=' * 100}\n{title}\n{'=' * 100}", flush=True)


def check(feature: str, ok: bool, detail: str, status: str | None = None) -> None:
    status = status or ("PASS" if ok else "FAIL")
    results.append((feature, status, detail))
    print(f"  >>> [{status}] {feature}: {detail}", flush=True)


def page_status(db_path: Path, start: int, end: int, label: str = "database status") -> list[tuple]:
    conn = sqlite3.connect(f"{db_path.resolve().as_uri()}?mode=ro", uri=True)
    rows = conn.execute("SELECT page_number, status, attempts, records_found FROM page_progress "
                        "WHERE page_number BETWEEN ? AND ? ORDER BY 1", (start, end)).fetchall()
    total = conn.execute("SELECT COUNT(*) FROM verified_records").fetchone()[0]
    conn.close()
    cells = "  ".join(f"{p}:{s}({a}x,{'-' if r is None else r} rec)" for p, s, a, r in rows)
    print(f"  [{label}] {cells} | verified records: {total}", flush=True)
    return rows


def stored_numbers(db_path: Path) -> list[int]:
    conn = sqlite3.connect(f"{db_path.resolve().as_uri()}?mode=ro", uri=True)
    numbers = [r[0] for r in conn.execute("SELECT number FROM verified_records ORDER BY number")]
    conn.close()
    return numbers


def digest(path: Path) -> str | None:
    return hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else None


# --- synthetic site behind the mocked browser ------------------------------------------------

class SyntheticSite(FakePage):
    """Mocked browser tab: serves a synthetic listing for whatever page is requested.

    ``events`` queues one-off behaviours per page: "interrupt" (Ctrl+C during the load),
    "incomplete" (a listing with missing rows) or an HTTP status code.
    """

    def __init__(self, events: dict[int, list[str | int]] | None = None) -> None:
        super().__init__()
        self.events = events or {}
        self.visited: list[int] = []

    def goto(self, url: str, wait_until: str = "domcontentloaded", timeout: int = 30_000) -> FakeResponse | None:
        if url == BLANK_URL:
            return super().goto(url, wait_until, timeout)
        page = int(parse_qs(urlsplit(url).query).get("page", ["1"])[0])
        self.calls.append(url)
        self.visited.append(page)
        event = self.events[page].pop(0) if self.events.get(page) else None
        if event == "interrupt":
            raise KeyboardInterrupt
        status, html = 200, listing_page(page, VERIFIED.get(page, []))
        if event == "incomplete":
            html = listing_html(listing_row(n) for n in list(page_segments(page))[:40])
        elif isinstance(event, int):
            status, html = event, f"<html><body>HTTP {event}</body></html>"
        self.url, self.html = url, html
        return FakeResponse(status)


class Browsers:
    """Gives each CLI invocation a fresh mocked browser session sharing one synthetic site."""

    def __init__(self, site: SyntheticSite) -> None:
        self.site = site
        self.sessions: list[FakePlaywright] = []
        self._original = browser.BrowserSession

    def __enter__(self) -> Browsers:
        def session(settings: BrowserSettings):
            fake = FakePlaywright()
            fake.browser.context.page = self.site
            self.sessions.append(fake)
            return self._original(settings, factory(fake))
        browser.BrowserSession = session
        return self

    def __exit__(self, *exc: object) -> None:
        browser.BrowserSession = self._original


def run_cli(*argv: str) -> int:
    print(f"\n$ python main.py {' '.join(argv)}", flush=True)
    code = cli.main(list(argv))
    print(f"(exit code {code})", flush=True)
    return code


class StatusAfterEachPage:
    """Prints the committed database state each time the runner reports a finished page."""

    def __init__(self, db_path: Path) -> None:
        self.db_path = db_path
        self._original = ProgressReporter.page_finished

    def __enter__(self) -> StatusAfterEachPage:
        original, db_path = self._original, self.db_path

        def page_finished(reporter: ProgressReporter, page: int, completed: bool) -> None:
            original(reporter, page, completed)
            page_status(db_path, reporter.start, reporter.end, f"saved after page {page}")

        ProgressReporter.page_finished = page_finished  # type: ignore[method-assign]
        return self

    def __exit__(self, *exc: object) -> None:
        ProgressReporter.page_finished = self._original  # type: ignore[method-assign]


# --- mocked-browser scenarios for CAPTCHA and errors -----------------------------------------

class VirtualClock:
    """Stands in for time, so retry back-offs and CAPTCHA waits run instantly."""

    def __init__(self, on_sleep=None) -> None:
        self.now, self.sleeps, self.on_sleep = 0.0, [], on_sleep

    def __call__(self) -> float:
        return self.now

    def sleep(self, seconds: float) -> None:
        self.sleeps.append(round(seconds, 1))
        self.now += seconds
        if self.on_sleep:
            self.on_sleep(len(self.sleeps))


def mocked_fetcher(page: FakePage, clock: VirtualClock, notes: list[str] | None = None,
                   **overrides: object) -> ListingFetcher:
    options = dict(page_delay_s=0, page_delay_jitter_s=0, captcha_poll_s=2, max_retries=2)
    options.update(overrides)
    notes = notes if notes is not None else []

    def notify(message: str) -> None:
        notes.append(message)
        print(f"  [scraper] {message}")

    return ListingFetcher(page, BrowserSettings(**options), notify=notify,  # type: ignore[arg-type]
                          sleep=clock.sleep, clock=clock, rng=random.Random(0))


def captcha_and_error_scenarios(tmp: Path) -> None:
    section("8. CAPTCHA AND ERROR HANDLING (mocked browser responses only)")

    print("\n-- 8a. CAPTCHA detection in --headless mode: the run stops at once, nothing is answered")
    page, clock = FakePage([captcha()]), VirtualClock()
    try:
        mocked_fetcher(page, clock, headless=True)(1)
        check("CAPTCHA detection", False, "no stop")
    except CaptchaRequiredError as exc:
        print(f"  CaptchaRequiredError: {exc}")
        check("CAPTCHA detection", len(page.calls) == 1, "detected on the first navigation, run stops")

    print("\n-- 8b. Visible browser: a (simulated) person answers wrongly twice, then correctly")
    page = FakePage([captcha("https://honestpool.ru/captcha?url=puzzle?id=71&page=2"), success(2, [150])])

    def person(poll: int) -> None:
        if poll <= 2:
            page.url = "https://honestpool.ru/captcha?url=puzzle?id=71&page=2"
            page.html = CAPTCHA_PAGE.replace("</body>", f"<img src='/captcha/img?{poll}'></body>")
        elif poll == 3:
            page.url, page.html = BASE, listing_page(1, [1, 2])

    notes: list[str] = []
    fetch = mocked_fetcher(page, VirtualClock(person), notes)
    html = fetch(2)
    rejected = sum("answer may have been rejected" in n for n in notes)
    check("Incorrect CAPTCHA answers keep waiting", rejected == 2,
          f"{rejected} rejected answers reported, kept waiting")
    check("Valid listing detected after manual verification", "tag_search_150" in html and
          fetch.manual_verifications == 1, "complete listing confirmed, page 2 then loaded in the same session")

    print("\n-- 8c. A second CAPTCHA in the same session stops the run safely (temporary database)")
    errors_db = tmp / "errors.db"
    page = FakePage([captcha(), success(2, [150]), captcha()])

    def person_once(poll: int) -> None:
        if poll == 1:
            page.url, page.html = BASE, listing_page(1)

    clock = VirtualClock(person_once)
    with Database(errors_db) as db:
        summary = run_pages(db, *plan_run(db, 2, 4), mocked_fetcher(page, clock))
    print(f"  run stopped: {summary.stopped}")
    rows = page_status(errors_db, 2, 4)
    check("Second CAPTCHA causes a safe stop", bool(summary.stopped) and
          [(s, a) for _, s, a, _ in rows] == [("completed", 1), ("pending", 0), ("pending", 0)],
          "page 2 kept, pages 3-4 pending with no attempt counted")

    print("\n-- 8d. HTTP 403 during a CLI run (full pipeline, mocked browser)")
    site = SyntheticSite({3001: [403]})
    with Browsers(site) as browsers:
        code = run_cli("--db", str(errors_db), "--start-page", "3000", "--end-page", "3002", "--live",
                       "--max-pages", "3", *PACING)
    rows = page_status(errors_db, 3000, 3002)
    launches = browsers.sessions[0].chromium.launches
    context_options = sorted(browsers.sessions[0].browser.context.kwargs)
    check("HTTP 403 stops safely", code == 3 and site.visited == [3000, 3001] and rows[1][1] == "pending",
          "exit 3 after page 3001, no retry, page 3002 never requested")
    check("No proxy rotation or relaunch on denial", launches == [{"headless": False}] and
          context_options == ["locale", "viewport"],
          f"1 browser launch {launches}, context options {context_options} only")

    print("\n-- 8e. HTTP 401")
    page = FakePage([http(401)])
    try:
        mocked_fetcher(page, VirtualClock())(1)
        check("HTTP 401 stops safely", False, "no stop")
    except AccessRestrictedError as exc:
        print(f"  AccessRestrictedError: {exc}")
        check("HTTP 401 stops safely", len(page.calls) == 1, "stopped after 1 navigation, no retry")

    print("\n-- 8f. HTTP 429 with Retry-After: 45, then success; persistent 429; Retry-After above the cap")
    clock = VirtualClock()
    html = mocked_fetcher(FakePage([http(429, headers={"Retry-After": "45"}), success(1, [1])]), clock)(1)
    print(f"  waits: {clock.sleeps} s")
    ok_retry_after = "tag_search_1" in html and clock.sleeps and clock.sleeps[0] >= 45
    page = FakePage([http(429)] * 5)
    try:
        mocked_fetcher(page, VirtualClock())(1)
        persistent = False
    except RateLimitedError as exc:
        print(f"  RateLimitedError: {exc}")
        persistent = len(page.calls) == 3
    page = FakePage([http(429, headers={"Retry-After": "3600"})])
    try:
        mocked_fetcher(page, VirtualClock(), max_retry_after_s=600)(1)
        capped = False
    except RateLimitedError as exc:
        print(f"  RateLimitedError: {exc}")
        capped = len(page.calls) == 1
    check("HTTP 429 and Retry-After", bool(ok_retry_after and persistent and capped),
          "waited 45 s then succeeded; persistent 429 stopped after 3 navigations; 3600 s request refused")

    print("\n-- 8g. HTTP 503 and timeouts: bounded retries with exponential back-off (--retries 2)")
    clock = VirtualClock()
    html = mocked_fetcher(FakePage([http(503), success(1, [1])]), clock, backoff_base_s=10)(1)
    print(f"  503 then success, waits: {clock.sleeps} s")
    page, clock = FakePage([timeout()] * 10), VirtualClock()
    try:
        mocked_fetcher(page, clock, backoff_base_s=10)(1)
        bounded = False
    except TransientFetchError as exc:
        print(f"  TransientFetchError: {exc}; back-off waits {clock.sleeps} s")
        bounded = len(page.calls) == 3
    check("HTTP 5xx and timeouts with bounded retries", "tag_search_1" in html and bounded,
          "503 retried once then succeeded; timeouts gave up after 1 + 2 retries")

    print("\n-- 8h. Automatic CAPTCHA solving")
    check("Automatic CAPTCHA solving", True,
          "by design the scraper only waits for a person and never answers the CAPTCHA", "NOT IMPLEMENTED")


# --- the walkthrough ---------------------------------------------------------------------------

def main() -> int:
    logging.basicConfig(stream=sys.stdout, level=logging.INFO, datefmt="%H:%M:%S",
                        format="%(asctime)s %(levelname)s %(name)s: %(message)s")
    before = {p: digest(p) for p in PROTECTED}
    tmp = Path(tempfile.mkdtemp(prefix="honestpool-demo-"))
    db_path = tmp / "demo.db"
    assert db_path.resolve() != DB_PATH.resolve()

    section(f"1. TEMPORARY DEMONSTRATION DATABASE: {db_path}")
    print("  The real database is not used. Pages come from a mocked browser serving synthetic HTML.")
    print(f"  Synthetic verified segments: {VERIFIED}")
    check("Temporary database", not db_path.exists(), "new empty location under /tmp")

    section("2-3. START AT PAGE 1000, END AT 1005; PROGRESS SAVED AFTER EACH PAGE")
    run_cli("--db", str(db_path), "--start-page", "1000", "--end-page", "1005")
    rows = page_status(db_path, 1, 10**6)
    check("Range selection (planning only)", [p for p, *_ in rows] == list(range(1000, 1006)) and
          {s for _, s, _, _ in rows} == {"pending"}, "pages 1000-1005 registered as pending, nothing fetched")
    site = SyntheticSite({1003: ["interrupt"]})
    with Browsers(site), StatusAfterEachPage(db_path):
        code = run_cli("--db", str(db_path), "--start-page", "1000", "--end-page", "1005", "--live",
                       "--max-pages", "6", *PACING)

    section("4. INTERRUPTION AND RECOVERY")
    print("  The run above was interrupted (Ctrl+C) while page 1003 was loading.")
    rows = page_status(db_path, 1000, 1005, "after interruption")
    check("Progress saved after each page", [s for _, s, _, _ in rows[:3]] == ["completed"] * 3,
          "pages 1000-1002 committed one by one before the interruption")
    check("Interruption is safe", code == 130 and rows[3][1:3] == ("pending", 0),
          "exit 130, page 1003 back to pending with no attempt counted")
    first_visits = list(site.visited)
    site.visited.clear()
    with Browsers(site), StatusAfterEachPage(db_path):
        code = run_cli("--db", str(db_path), "--resume", "--live", "--max-pages", "6", *PACING)
    rows = page_status(db_path, 1000, 1005, "after resume")
    numbers = stored_numbers(db_path)
    expected = sorted(n for p in range(1000, 1006) for n in VERIFIED[p])
    check("Resume skips completed pages", code == 0 and site.visited == [1003, 1004, 1005],
          f"resume requested only {site.visited}")
    check("No duplicates or lost records", numbers == expected and len(numbers) == len(set(numbers)),
          f"{len(numbers)} unique records {numbers}")
    check("Only pages inside 1000-1005 processed",
          set(first_visits + site.visited) <= set(range(1000, 1006)) and len(rows) == 6,
          f"pages requested: {first_visits} then {site.visited}")

    section("5. CHANGING THE STARTING PAGE: RANGE 2000-2003 (page 2001 first returns an incomplete listing)")
    site = SyntheticSite({2001: ["incomplete"]})
    with Browsers(site), StatusAfterEachPage(db_path):
        run_cli("--db", str(db_path), "--start-page", "2000", "--end-page", "2003", "--live",
                "--max-pages", "4", *PACING)
    rows_old = page_status(db_path, 1000, 1005, "range 1000-1005")
    rows_new = page_status(db_path, 2000, 2003, "range 2000-2003")
    numbers_after = stored_numbers(db_path)
    check("New range processed", site.visited == [2000, 2001, 2002, 2003],
          "pages 2000-2003 requested; page 2001 failed validation and saved nothing")
    check("Earlier data intact", [s for _, s, _, _ in rows_old] == ["completed"] * 6 and
          set(expected) <= set(numbers_after), "range 1000-1005 still completed with its 6 records")
    check("Invalid listing marked failed", rows_new[1][1] == "failed", "page 2001 failed (1 attempt)")

    section("7. PROGRESS REPORTING")
    print("  a) Resume range 2000-2003: the report before the page shows 1 failed page.")
    site.visited.clear()
    with Browsers(site):
        run_cli("--db", str(db_path), "--resume", "--live", "--max-pages", "1", *PACING)
    rows_new = page_status(db_path, 2000, 2003, "range 2000-2003")
    check("Failed page retried on resume", site.visited == [2001] and rows_new[1][1] == "completed",
          "page 2001 retried first and completed")
    print("\n  b) Extend the range to 1000-1100 and fetch 3 pages: the estimate uses the measured pace.")
    site.visited.clear()
    with Browsers(site):
        run_cli("--db", str(db_path), "--start-page", "1000", "--end-page", "1100", "--live",
                "--max-pages", "3", *PACING)
    with Database(db_path, readonly=True) as db:
        final = ProgressReporter(db, 1000, 1100, notify=lambda _: None).snapshot()
    check("Progress reporting", site.visited == [1006, 1007, 1008] and final.completed == 9 and
          final.pending == 92 and final.failed == 0,
          "current page, completed/pending/failed, verified records and approximate remaining time shown")

    section("6. CSV AND XLSX EXPORT OF THE TEMPORARY DATABASE")
    csv_path, xlsx_path = tmp / "verified_records.csv", tmp / "verified_records.xlsx"
    run_cli("--db", str(db_path), "--export", "csv", "--output", str(csv_path))
    run_cli("--db", str(db_path), "--export", "xlsx", "--output", str(xlsx_path))
    with csv_path.open(encoding="utf-8-sig", newline="") as fh:
        csv_rows = list(csv.reader(fh))
    print(f"\n  {csv_path.name}:")
    for r in csv_rows:
        print("    " + ",".join(r))
    ws = load_workbook(xlsx_path).active
    xlsx_rows = [[c.value for c in r] for r in ws.iter_rows()]
    hex_cells = [c for r in list(ws.iter_rows())[1:] for c in r[1:3]]
    inline_cells = zipfile.ZipFile(xlsx_path).read("xl/worksheets/sheet1.xml").decode().count('t="inlineStr"')
    print(f"\n  {xlsx_path.name}: {len(xlsx_rows) - 1} rows; hex cell types "
          f"{sorted({c.data_type for c in hex_cells})}, number formats {sorted({c.number_format for c in hex_cells})}, "
          f"inline-string cells in the sheet XML: {inline_cells}")
    for r in xlsx_rows:
        print("    " + " | ".join(str(v) for v in r))
    header = ["Number", "Start hex", "End hex", "Status"]
    stored = stored_numbers(db_path)
    check("Four required columns", csv_rows[0] == header and xlsx_rows[0] == header, ", ".join(header))
    check("Hex values complete", all((r[1], r[2]) == segment_range(int(r[0])) for r in csv_rows[1:]) and
          [int(r[0]) for r in csv_rows[1:]] == stored, f"{len(stored)} rows, 18-digit values matching the formula")
    check("XLSX stores hex as text", all(c.data_type == "s" and c.number_format == "@" for c in hex_cells) and
          [r[1:3] for r in xlsx_rows[1:]] == [r[1:3] for r in csv_rows[1:]], "string cells with Text format '@'")

    captcha_and_error_scenarios(tmp)

    section("REAL DATA PROTECTION")
    after = {p: digest(p) for p in PROTECTED}
    for p in PROTECTED:
        print(f"  {p.relative_to(PROJECT_ROOT)}: {'unchanged' if before[p] == after[p] else 'CHANGED'}")
    check("Real database and delivery package untouched", before == after, "checksums identical before and after")

    section("SUMMARY")
    width = max(len(f) for f, _, _ in results)
    for feature, status, _ in results:
        print(f"  {feature:<{width}}  {status}")
    failed = [f for f, s, _ in results if s == "FAIL"]
    print(f"\n  {len(results)} checks, {len(failed)} failed. Temporary files: {tmp}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
