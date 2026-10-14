"""End to end, offline: 10 pages through the real CLI, runner and SQLite layer, an interruption on
page 6, resume with a CAPTCHA solved through the mocked 2Captcha API, then CSV/XLSX export."""

from __future__ import annotations

import base64
import csv
import json
import logging
import os
import socket
import subprocess
import sys
import tempfile
import time
from collections.abc import Iterator
from pathlib import Path

import pytest
from openpyxl import load_workbook

import main as cli
from scraper import browser, two_captcha
from scraper.browser import CAPTCHA_INPUT_SELECTOR, BrowserSession
from scraper.config import DB_PATH, DEFAULT_BASE_URL, PROJECT_ROOT
from scraper.database import COMPLETED, FAILED, PENDING, PROCESSING, Database
from scraper.pagination import page_url
from scraper.parser import VERIFIED_TEXT
from scraper.runner import plan_run, run_pages
from tests.e2e_scenario import (
    DECOYS,
    INTERRUPTED_PAGE,
    PAGES,
    SNAPSHOT,
    VERIFIED,
    fast,
    row,
    scenario_page,
)
from tests.fake_playwright import CAPTCHA_PNG, FakePlaywright, ScriptedNav, captcha_form, factory
from tests.fake_two_captcha import FAKE_KEY, FakeTwoCaptchaAPI
from tests.html_fixtures import FakeSite, listing_html
from tests.platform_support import KILLED_RETURNCODE, open_after_crash, utf8_child_env

ANSWER = "k7m3q"
CAPTCHA_PAGE = 8
BEFORE = range(1, INTERRUPTED_PAGE)
AFTER = range(INTERRUPTED_PAGE, PAGES[-1] + 1)
ALL_VERIFIED = sorted(n for numbers in VERIFIED.values() for n in numbers)


def expected_range(number: int) -> tuple[str, str]:
    """The published layout, written out independently of scraper.parser.segment_range."""
    start = 0x400000000000000000 + (number - 1) * 2**44
    return format(start, "018x"), format(start + 2**44, "018x")


@pytest.fixture(autouse=True)
def no_network(monkeypatch: pytest.MonkeyPatch) -> None:
    def refuse(*args, **kwargs):
        raise AssertionError("network access attempted")

    monkeypatch.setattr(socket.socket, "connect", refuse)
    monkeypatch.setattr(socket, "create_connection", refuse)


@pytest.fixture
def workdir() -> Iterator[Path]:
    with tempfile.TemporaryDirectory(prefix="step21-") as name:
        yield Path(name)


def interrupted_run(db: Path, mode: str) -> tuple[int, list[dict]]:
    env = {k: v for k, v in utf8_child_env().items() if not k.startswith("CAPTCHA_")}
    env.update(PYTHONPATH=str(PROJECT_ROOT), PYTHONUNBUFFERED="1")
    done = subprocess.run([sys.executable, "-m", "tests.e2e_scenario", str(db), mode], cwd=PROJECT_ROOT, env=env,
                          capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120)
    snapshots = [json.loads(line[len(SNAPSHOT):]) for line in done.stdout.splitlines() if line.startswith(SNAPSHOT)]
    return done.returncode, snapshots


def auto_csv_numbers(directory: Path) -> list[int]:
    """Numbers in the CSV that live runs rewrite next to the database after every completed page."""
    with (directory / "verified_records.csv").open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        assert next(reader) == ["Number", "Start hex", "End hex", "Status"]
        return [int(row[0]) for row in reader]


def records_by_page(db: Database) -> dict[int, list[int]]:
    rows = db.conn.execute("SELECT number, source_page FROM verified_records ORDER BY number").fetchall()
    return {p: [n for n, src in rows if src == p] for p in PAGES}


def saved_timestamps(db: Database) -> dict[int, str]:
    return dict(db.conn.execute("SELECT number, extracted_at FROM verified_records").fetchall())


def test_scenario_pages_mix_verified_rows_with_decoy_statuses() -> None:
    assert sum(len(v) for v in VERIFIED.values()) == len(ALL_VERIFIED) == 14
    html = "".join(scenario_page(p) for p in PAGES)
    assert html.count(">Проверено</span>") == 14
    for colour, text in DECOYS:
        assert f'<span class="badge {colour}">{text}</span>' in html


@pytest.mark.parametrize("mode", ["kill-fetch", "kill-commit", "ctrl-c"])
def test_ten_pages_interruption_resume_captcha_and_export(
        mode: str, workdir: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str],
        caplog: pytest.LogCaptureFixture) -> None:
    db_path = workdir / "step21_test.db"
    assert db_path.resolve() != DB_PATH.resolve()

    # --- run 1: pages 1..10 in a separate process, stopped while page 6 is in progress --------------
    returncode, snapshots = interrupted_run(db_path, mode)
    assert returncode == (130 if mode == "ctrl-c" else KILLED_RETURNCODE)

    # Before each navigation, every earlier page is already committed with its records and checkpoint.
    assert [s["navigating_to"] for s in snapshots] == list(range(1, INTERRUPTED_PAGE + 1))
    for snap in snapshots:
        current = snap["navigating_to"]
        statuses = {p: (status, attempts) for p, status, attempts in snap["pages"]}
        assert all(statuses[p] == (COMPLETED, 1) for p in range(1, current))
        assert statuses[current] == (PROCESSING, 1)
        assert all(statuses[p] == (PENDING, 0) for p in range(current + 1, PAGES[-1] + 1))
        assert snap["records"] == [[n, p] for p in range(1, current) for n in VERIFIED[p]]

    with open_after_crash(db_path) as db:
        page6 = db.page(INTERRUPTED_PAGE)
        if mode == "ctrl-c":
            assert (page6["status"], page6["attempts"]) == (PENDING, 0)
            assert page6["last_error"] == "stopped before completion"
        else:  # killed: nothing ran on the way out, the page is still marked processing
            assert (page6["status"], page6["attempts"]) == (PROCESSING, 1)
        for p in BEFORE:
            assert (db.page(p)["status"], db.page(p)["attempts"], db.page(p)["records_found"]) == \
                   (COMPLETED, 1, len(VERIFIED[p]))
        untouched = [p for p in AFTER if p != INTERRUPTED_PAGE]
        assert all((db.page(p)["status"], db.page(p)["attempts"]) == (PENDING, 0) for p in untouched)
        by_page = records_by_page(db)
        assert {p: by_page[p] for p in BEFORE} == {p: VERIFIED[p] for p in BEFORE}
        assert all(by_page[p] == [] for p in AFTER)  # page 6's uncommitted inserts were rolled back
        assert db.record_count() == sum(len(VERIFIED[p]) for p in BEFORE) == 7
        assert db.load_run_range() == (1, 10)
        before_resume = saved_timestamps(db)
    assert auto_csv_numbers(workdir) == sorted(n for p in BEFORE for n in VERIFIED[p])

    # --- run 2: --resume with a CAPTCHA on page 8, solved through the mocked 2Captcha API ------------
    caplog.set_level(logging.DEBUG)
    fake = FakePlaywright()
    page = fake.browser.context.page
    nav = {p: ScriptedNav(html=scenario_page(p)) for p in AFTER}
    page.script = [nav[6], nav[7], captcha_form(1), nav[8], nav[9], nav[10]]
    page.submissions = [ScriptedNav(url=DEFAULT_BASE_URL, html=scenario_page(1))]  # redirect after the answer
    api = FakeTwoCaptchaAPI([ANSWER])
    monkeypatch.setattr(browser, "BrowserSession", lambda settings: BrowserSession(settings, factory(fake)))
    monkeypatch.setattr(two_captcha, "urllib_transport", api)
    monkeypatch.setattr(time, "sleep", lambda seconds: None)
    monkeypatch.setattr(cli, "run_live", fast(cli.run_live))
    monkeypatch.setenv("CAPTCHA_PROVIDER", "2captcha")
    monkeypatch.setenv("CAPTCHA_API_KEY", FAKE_KEY)
    monkeypatch.setenv("CAPTCHA_TIMEOUT", "120")

    resume = ["--db", str(db_path), "--resume", "--live", "--headless", "--max-pages", "10"]
    assert cli.main(resume) == 0
    out, err = capsys.readouterr()
    recovered = 0 if mode == "ctrl-c" else 1
    assert f"Recovered interrupted pages: {recovered}" in out
    assert "Pages to process: 5 (next: [6, 7, 8, 9, 10])" in out
    assert "Live run: 5 completed, 0 failed, 7 new verified records" in out
    assert "Page requests sent: 7" in out  # 6 navigations and 1 CAPTCHA answer submission
    assert "Verification succeeded automatically" in out

    urls = {p: page_url(DEFAULT_BASE_URL, p) for p in PAGES}
    assert page.calls == [urls[6], urls[7], urls[8], urls[8], urls[9], urls[10]]  # no completed page refetched
    assert api.methods == ["createTask", "getTaskResult", "reportCorrect"]
    assert page.screenshots and page.fills == [(CAPTCHA_INPUT_SELECTOR, ANSWER)]
    assert api.requests[0][1]["task"]["type"] == "ImageToTextTask"
    assert base64.b64decode(api.requests[0][1]["task"]["body"]) == CAPTCHA_PNG
    assert FAKE_KEY not in out + err + caplog.text and ANSWER not in out + err + caplog.text

    with Database(db_path) as db:
        for p in PAGES:
            info = db.page(p)
            attempts = 2 if p == INTERRUPTED_PAGE and mode != "ctrl-c" else 1
            assert (info["status"], info["attempts"], info["records_found"], info["last_error"]) == \
                   (COMPLETED, attempts, len(VERIFIED[p]), None)
        assert records_by_page(db) == VERIFIED
        after_resume = saved_timestamps(db)
        assert {n: after_resume[n] for n in before_resume} == before_resume  # earlier records untouched
    assert auto_csv_numbers(workdir) == ALL_VERIFIED

    # --- run 3: nothing is fetched twice, whether resuming or selecting the same range again -----------
    for argv in (resume, ["--db", str(db_path), "--start-page", "1", "--end-page", "10", "--live", "--headless",
                          "--max-pages", "10"]):
        assert cli.main(argv) == 0
        out = capsys.readouterr().out
        assert "Pages to process: 0" in out and "Verified records stored: 14" in out
    assert len(page.calls) == 6 and len(api.requests) == 3

    with Database(db_path, readonly=True) as db:
        rows = db.conn.execute("SELECT number, start_hex, end_hex, status, source_page FROM verified_records "
                               "ORDER BY number").fetchall()
        duplicates = db.conn.execute("SELECT number FROM verified_records GROUP BY number HAVING COUNT(*) > 1")
        assert duplicates.fetchall() == []
        assert db.progress_counts() == {PENDING: 0, PROCESSING: 0, COMPLETED: 10, FAILED: 0}
    assert [r[0] for r in rows] == ALL_VERIFIED
    for number, start_hex, end_hex, status, source_page in rows:
        assert (start_hex, end_hex) == expected_range(number) and len(start_hex) == len(end_hex) == 18
        assert status == VERIFIED_TEXT == "Проверено"
        assert number in VERIFIED[source_page]

    # --- export to new files, then compare them with SQLite ------------------------------------------
    csv_path, xlsx_path = workdir / "step21_verified_records.csv", workdir / "step21_verified_records.xlsx"
    for fmt, path in (("csv", csv_path), ("xlsx", xlsx_path)):
        assert cli.main(["--db", str(db_path), "--export", fmt, "--output", str(path)]) == 0
        assert f"Exported 14 verified records to {path}" in capsys.readouterr().out
    expected_rows = [(n, s, e, st) for n, s, e, st, _ in rows]

    with csv_path.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        assert next(reader) == ["Number", "Start hex", "End hex", "Status"]
        csv_rows = [(int(n), s, e, st) for n, s, e, st in reader]
    assert csv_rows == expected_rows
    assert (workdir / "verified_records.csv").read_bytes() == csv_path.read_bytes()
    assert list(workdir.glob(".verified_records.*")) == []

    sheet = load_workbook(xlsx_path).active
    assert [c.value for c in sheet[1]] == ["Number", "Start hex", "End hex", "Status"]
    xlsx_rows = list(sheet.iter_rows(min_row=2))
    assert [(a.value, b.value, c.value, d.value) for a, b, c, d in xlsx_rows] == expected_rows
    for number_cell, start_cell, end_cell, _ in xlsx_rows:
        assert number_cell.data_type == "n"
        for cell in (start_cell, end_cell):
            assert cell.data_type == "s" and cell.number_format == "@" and len(cell.value) == 18


def test_hard_kill_while_rewriting_the_csv_keeps_the_previous_csv_and_resume_repairs_it(
        workdir: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]) -> None:
    db_path = workdir / "step21_test.db"
    returncode, snapshots = interrupted_run(db_path, "kill-csv")
    assert returncode == KILLED_RETURNCODE
    assert [s["navigating_to"] for s in snapshots] == list(range(1, INTERRUPTED_PAGE + 1))

    with open_after_crash(db_path) as db:
        assert (db.page(INTERRUPTED_PAGE)["status"], db.page(INTERRUPTED_PAGE)["attempts"]) == (COMPLETED, 1)
        assert records_by_page(db) == {p: VERIFIED[p] if p <= INTERRUPTED_PAGE else [] for p in PAGES}
    # SQLite already holds page 6; the CSV is still the complete file written after page 5.
    assert auto_csv_numbers(workdir) == sorted(n for p in BEFORE for n in VERIFIED[p])
    leftovers = list(workdir.glob(".verified_records.*.csv"))
    assert len(leftovers) == 1, "the killed process left its unfinished temporary file"
    old = time.time() - 3600
    os.utime(leftovers[0], (old, old))

    fake = FakePlaywright()
    fake.browser.context.page.script = [ScriptedNav(html=scenario_page(p)) for p in range(INTERRUPTED_PAGE + 1, 11)]
    monkeypatch.setattr(browser, "BrowserSession", lambda settings: BrowserSession(settings, factory(fake)))
    monkeypatch.setattr(time, "sleep", lambda seconds: None)
    monkeypatch.setattr(cli, "run_live", fast(cli.run_live))
    for name in ("CAPTCHA_PROVIDER", "CAPTCHA_API_KEY", "CAPTCHA_TIMEOUT"):
        monkeypatch.delenv(name, raising=False)
    assert cli.main(["--db", str(db_path), "--resume", "--live", "--headless", "--max-pages", "10"]) == 0
    out = capsys.readouterr().out
    assert "Pages to process: 4 (next: [7, 8, 9, 10])" in out
    assert auto_csv_numbers(workdir) == ALL_VERIFIED
    assert list(workdir.glob(".verified_records.*")) == []


def test_verified_text_with_a_non_green_badge_saves_nothing_from_that_page(tmp_path: Path) -> None:
    grey = listing_html(row(n, "bg-secondary", "Проверено") if n == 150 else row(n, "bg-success", "Проверено")
                        if n == 101 else row(n, *DECOYS[0]) for n in range(101, 201))
    site = FakeSite(verified={1: [9], 3: [250]}, scripted={2: [grey]})
    with Database(tmp_path / "badge.db") as db:
        summary = run_pages(db, *plan_run(db, 1, 3), site)
        assert (summary.completed, summary.failed) == (2, 1)
        assert db.page(2)["status"] == FAILED and "wrong_badge" in db.page(2)["last_error"]
        assert [r.number for r in db.iter_records()] == [9, 250]  # not even the green 101 from page 2
