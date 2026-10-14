"""CSV and XLSX export tests, exporting directly from SQLite."""

from __future__ import annotations

import csv
import re
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scraper import exporter
from scraper.database import Database
from scraper.exporter import CSV_COLUMNS, export_csv, export_from_db, export_records, export_xlsx
from tests.html_fixtures import record

# Digit-only values are the ones Excel would otherwise turn into 4.00001E+17.
NUMBERS = [1, 2, 10, 101, 838640]


@pytest.fixture
def filled_db(db: Database) -> Database:
    for page in sorted({(n - 1) // 100 + 1 for n in NUMBERS}):
        db.claim_page(page)
        db.complete_page(page, [record(n) for n in NUMBERS if (n - 1) // 100 + 1 == page], 100)
    return db


def test_csv_export_from_database(filled_db: Database, tmp_path: Path) -> None:
    out = tmp_path / "verified.csv"
    assert export_from_db(filled_db, out) == len(NUMBERS)
    raw = out.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    with out.open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    assert tuple(rows[0]) == CSV_COLUMNS
    assert rows[1:] == [[str(r.number), r.start_hex, r.end_hex, r.status] for r in map(record, NUMBERS)]


def test_xlsx_keeps_hex_as_exact_text(filled_db: Database, tmp_path: Path) -> None:
    out = tmp_path / "verified.xlsx"
    assert export_from_db(filled_db, out) == len(NUMBERS)
    ws = load_workbook(out).active
    assert [c.value for c in ws[1]] == list(CSV_COLUMNS)
    for row, number in zip(ws.iter_rows(min_row=2), NUMBERS, strict=True):
        expected = record(number)
        assert row[0].value == number and row[0].data_type == "n"
        for cell, value in ((row[1], expected.start_hex), (row[2], expected.end_hex)):
            assert cell.value == value and isinstance(cell.value, str) and len(cell.value) == 18
            assert cell.data_type == "s"
            assert cell.number_format == "@"
        assert row[3].value == "Проверено"
    assert ws.max_row == len(NUMBERS) + 1


def test_xlsx_stores_hex_as_strings_in_sheet_xml(filled_db: Database, tmp_path: Path) -> None:
    out = tmp_path / "verified.xlsx"
    export_from_db(filled_db, out)
    with zipfile.ZipFile(out) as zf:
        sheet = zf.read("xl/worksheets/sheet1.xml").decode()
    hex_cells = re.findall(r'<c r="[BC]\d+"[^>]*>', sheet)
    assert len(hex_cells) == 2 * (len(NUMBERS) + 1)
    # Both string encodings are text to Excel; a numeric cell would have no t= or t="n".
    assert all(re.search(r't="(s|inlineStr|str)"', c) for c in hex_cells)
    assert '<c r="B3" s="3" t="inlineStr"><is><t>400000100000000000</t></is></c>' in sheet
    assert "E+" not in sheet


def test_xlsx_layout_is_readable(filled_db: Database, tmp_path: Path) -> None:
    out = tmp_path / "verified.xlsx"
    export_from_db(filled_db, out)
    ws = load_workbook(out).active
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == f"A1:D{len(NUMBERS) + 1}"
    assert all(ws[f"{col}1"].font.bold for col in "ABCD")
    assert ws.column_dimensions["B"].width >= 20 and ws.column_dimensions["C"].width >= 20


def test_empty_database_exports_header_only(db: Database, tmp_path: Path) -> None:
    assert export_from_db(db, tmp_path / "empty.xlsx") == 0
    assert export_from_db(db, tmp_path / "empty.csv") == 0
    assert load_workbook(tmp_path / "empty.xlsx").active.max_row == 1


def test_export_records_dispatches_by_suffix(tmp_path: Path) -> None:
    records = [record(2), record(1)]
    assert export_records(records, tmp_path / "a.CSV") == 2
    assert export_records(records, tmp_path / "a.xlsx") == 2
    with pytest.raises(ValueError):
        export_records(records, tmp_path / "a.json")


def test_failed_export_leaves_existing_file_untouched(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    out = tmp_path / "verified.xlsx"
    export_xlsx([record(1)], out)
    before = out.read_bytes()

    def broken_save(self: object, filename: object) -> None:
        Path(filename).write_bytes(b"partial")
        raise OSError("disk full")

    monkeypatch.setattr(exporter.Workbook, "save", broken_save)
    with pytest.raises(OSError):
        export_xlsx([record(1), record(2)], out)
    assert out.read_bytes() == before
    assert [p.name for p in tmp_path.iterdir()] == ["verified.xlsx"]


def test_csv_export_keeps_step4_format(tmp_path: Path) -> None:
    out = tmp_path / "v.csv"
    export_csv([record(10), record(1)], out)
    assert out.read_text(encoding="utf-8-sig").splitlines() == [
        "Number,Start hex,End hex,Status",
        "1,400000000000000000,400000100000000000,Проверено",
        "10,400000900000000000,400000a00000000000,Проверено",
    ]
