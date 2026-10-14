"""Offline import of verified segment numbers (synthetic fixtures, temporary databases only)."""

from __future__ import annotations

import csv
import hashlib
from pathlib import Path

import pytest
from openpyxl import load_workbook

import main as cli
from scraper.config import PROJECT_ROOT, SEGMENT_COUNT
from scraper.database import IMPORTED, WEBSITE, Database, InvalidRecordError
from scraper.exporter import export_from_db
from scraper.importer import ImportFileError, import_file, parse_segment_number, read_verified_file
from scraper.parser import VERIFIED_TEXT, SegmentRecord, segment_range
from scraper.runner import plan_run, process_page, run_pages
from tests.html_fixtures import FakeSite

FIXTURES = Path(__file__).parent / "fixtures" / "import"
SOURCE = "synthetic test list"


def write(tmp_path: Path, name: str, content: str | bytes) -> Path:
    path = tmp_path / name
    if isinstance(content, bytes):
        path.write_bytes(content)
    else:
        path.write_text(content, encoding="utf-8")
    return path


def page_rows(db: Database) -> list[tuple]:
    return [tuple(r) for r in db.conn.execute("SELECT * FROM page_progress ORDER BY page_number")]


def website_rows(db: Database) -> list[tuple]:
    return [tuple(r) for r in db.conn.execute("SELECT * FROM verified_records ORDER BY number")]


def seed_website_page(db: Database) -> None:
    """Page 1 scraped from the (fake, offline) site with verified numbers 1 and 2; pages 2-3 pending."""
    plan_run(db, 1, 3)
    process_page(db, 1, FakeSite(verified={1: [1, 2]}))


# --- PART B: hexadecimal calculation ---------------------------------------------

@pytest.mark.parametrize(("number", "start", "end"), [
    (1, "400000000000000000", "400000100000000000"),
    (2, "400000100000000000", "400000200000000000"),
    (150, "400009500000000000", "400009600000000000"),
    (250, "40000f900000000000", "40000fa00000000000"),
    (SEGMENT_COUNT, "7ffffff00000000000", "800000000000000000"),
])
def test_hex_range_matches_the_specified_formula(tmp_path: Path, number: int, start: str, end: str) -> None:
    computed = 0x400000000000000000 + (number - 1) * 2**44
    assert (start, end) == (f"{computed:018x}", f"{computed + 2**44:018x}")
    assert segment_range(number) == (start, end)
    record = read_verified_file(write(tmp_path, "one.txt", f"{number}\n")).records()[0]
    assert record == SegmentRecord(number, start, end, VERIFIED_TEXT) and record.status == "Проверено"
    assert all(len(h) == 18 and h == h.lower() and not h.startswith("0x") for h in (start, end))


# --- PART A/C: valid imports -----------------------------------------------------

def test_valid_txt_import(db: Database) -> None:
    verified, outcome = import_file(db, FIXTURES / "verified_numbers.txt", SOURCE)
    assert verified.numbers == (150, 250, SEGMENT_COUNT) and verified.duplicates == 0
    assert (outcome.inserted, outcome.already_scraped, outcome.already_imported) == (3, 0, 0)
    assert [r.number for r in db.iter_imported_records()] == [150, 250, SEGMENT_COUNT]
    assert db.imported_row(250)["start_hex"] == "40000f900000000000"


@pytest.mark.parametrize(("name", "numbers"), [("verified_numbers.csv", (150, 250)),
                                               ("numbers_only.csv", (2, SEGMENT_COUNT))])
def test_valid_csv_import(db: Database, name: str, numbers: tuple[int, ...]) -> None:
    verified, outcome = import_file(db, FIXTURES / name, SOURCE)
    assert verified.numbers == numbers and outcome.inserted == len(numbers)
    assert tuple(r.number for r in db.iter_imported_records()) == numbers


def test_txt_accepts_crlf_bom_and_surrounding_spaces(tmp_path: Path) -> None:
    path = write(tmp_path, "list.txt", "\ufeff 12 \r\n\r\n13\r\n")
    assert read_verified_file(path).numbers == (12, 13)


def test_cli_import_with_stated_source(db_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    code = cli.main(["--db", str(db_path), "--import-verified", str(FIXTURES / "verified_numbers.txt"),
                     "--source", SOURCE])
    out = capsys.readouterr().out
    assert code == 0
    assert "Segment numbers: 3 valid, 0 duplicate(s) removed, 3 unique" in out
    assert f"Imported 3 new verified record(s) from {SOURCE!r} (import batch 1)" in out
    assert "Website page progress was not changed" in out
    assert "Verified records stored: 0 from the website, 3 imported" in out


def test_cli_asks_for_the_source_interactively(db_path: Path, monkeypatch: pytest.MonkeyPatch,
                                               capsys: pytest.CaptureFixture[str]) -> None:
    prompts: list[str] = []
    monkeypatch.setattr(cli, "ask_source", lambda prompt: prompts.append(prompt) or "operator export, Oct 2026")
    assert cli.main(["--db", str(db_path), "--import-verified", str(FIXTURES / "verified_numbers.csv")]) == 0
    assert "trusted, authorized source" in capsys.readouterr().out and "authorized source" in prompts[0]
    with Database(db_path) as db:
        assert db.import_batches()[0]["source"] == "operator export, Oct 2026"


@pytest.mark.parametrize("answer", [None, "", "   "])
def test_cli_refuses_without_a_stated_source(db_path: Path, monkeypatch: pytest.MonkeyPatch,
                                             capsys: pytest.CaptureFixture[str], answer: str | None) -> None:
    monkeypatch.setattr(cli, "ask_source", lambda prompt: answer)
    assert cli.main(["--db", str(db_path), "--import-verified", str(FIXTURES / "verified_numbers.txt")]) == 4
    assert "import refused, nothing was written" in capsys.readouterr().err
    assert not db_path.exists()


def test_non_interactive_import_without_source_is_refused(db_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(cli.sys.stdin, "isatty", lambda: False, raising=False)
    assert cli.ask_source("?") is None
    assert cli.main(["--db", str(db_path), "--import-verified", str(FIXTURES / "verified_numbers.txt")]) == 4
    assert not db_path.exists()


# --- PART C: invalid values, duplicates, empty and malformed files -----------------

@pytest.mark.parametrize("value", ["0", "-1", "+5", "1.0", "1e3", "0x10", "1,000", "1 000", "67108865",
                                   "abc", "\u0663", "12\u200b", "99999999999999999999999"])
def test_invalid_segment_numbers_are_rejected(value: str) -> None:
    with pytest.raises(ValueError):
        parse_segment_number(value)


def test_one_invalid_value_rejects_the_whole_file(db: Database) -> None:
    with pytest.raises(ImportFileError) as exc:
        import_file(db, FIXTURES / "invalid_numbers.txt", SOURCE)
    assert [p.split(":")[0] for p in exc.value.problems] == ["line 2", "line 3", "line 4", "line 5"]
    assert "outside 1..67,108,864" in str(exc.value)
    assert db.imported_count() == 0 and db.import_batches() == [] and not db.has_imports()


def test_cli_invalid_file_writes_nothing(db_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    assert cli.main(["--db", str(db_path), "--import-verified", str(FIXTURES / "invalid_numbers.txt"),
                     "--source", SOURCE]) == 4
    err = capsys.readouterr().err
    assert "import refused, nothing was written: invalid_numbers.txt: line 2: 'abc' is not a whole" in err
    assert not db_path.exists()


def test_duplicates_are_removed(db: Database) -> None:
    verified, outcome = import_file(db, FIXTURES / "duplicates.txt", SOURCE)
    assert (verified.values_read, verified.numbers, verified.duplicates) == (5, (5, 7), 3)
    assert outcome.inserted == 2 and db.import_batches()[0]["duplicates"] == 3


@pytest.mark.parametrize(("name", "content"), [
    ("empty.txt", ""), ("blank.txt", "\n  \n\r\n"), ("bom.txt", "\ufeff"),
    ("empty.csv", ""), ("header_only.csv", "Number,Status\n"),
])
def test_empty_files_are_rejected(tmp_path: Path, name: str, content: str) -> None:
    path = FIXTURES / name if name == "empty.txt" else write(tmp_path, name, content)
    with pytest.raises(ImportFileError, match="no segment numbers found|no column named"):
        read_verified_file(path)


@pytest.mark.parametrize(("content", "message"), [
    ("Number,Comment\n150,\"unterminated\n250,ok\n", "malformed CSV"),
    ("Segment,Status\n150,Проверено\n", "no column named 'Number'"),
    ("Number;Status\n150;Проверено\n", "no column named 'Number'"),
    ("Number,Number\n1,2\n", "repeated column"),
    ("Number,Status\n150\n", r"line 2: 1 field\(s\), the header has 2"),
    ("Number,Status\n150,Проверено,extra\n", r"3 field\(s\), the header has 2"),
    ("Number,Start hex,End hex\n150,400009500000000000,400009700000000000\n", "End hex"),
    ("Number,Start hex\n150,0x400009500000000000\n", "Start hex"),
    ("Number,Status\n150,Не проверено\n", "only verified numbers"),
    ("Number,Status\n150,\n", "only verified numbers"),
    ("Number\n\n,\n", r"line 3: 2 field\(s\), the header has 1"),
    ("Number,Comment\n,note\n", "line 2: Number empty value"),
])
def test_malformed_csv_is_rejected(tmp_path: Path, content: str, message: str) -> None:
    with pytest.raises(ImportFileError, match=message):
        read_verified_file(write(tmp_path, "bad.csv", content))


def test_fixture_malformed_csvs_are_rejected() -> None:
    for name in ("malformed_quote.csv", "missing_number_column.csv"):
        with pytest.raises(ImportFileError):
            read_verified_file(FIXTURES / name)


@pytest.mark.parametrize(("name", "content", "message"), [
    ("list.xlsx", b"PK\x03\x04", "unsupported file type"),
    ("list", b"1\n", "unsupported file type"),
    ("latin1.txt", "150\nN\u00famero\n".encode("latin-1"), "not UTF-8"),
    ("binary.txt", b"150\n\x00\x01\n", "NUL bytes"),
])
def test_unsupported_or_binary_files_are_rejected(tmp_path: Path, name: str, content: bytes, message: str) -> None:
    with pytest.raises(ImportFileError, match=message):
        read_verified_file(write(tmp_path, name, content))


def test_missing_file_is_rejected(tmp_path: Path) -> None:
    with pytest.raises(ImportFileError, match="cannot read file"):
        read_verified_file(tmp_path / "absent.txt")


def test_database_layer_rejects_wrong_values(db: Database) -> None:
    good = SegmentRecord(150, *segment_range(150), VERIFIED_TEXT)
    for bad in (SegmentRecord(150, "0" * 18, "1" * 18, VERIFIED_TEXT),
                SegmentRecord(150, *segment_range(150), "Не проверено")):
        with pytest.raises(InvalidRecordError):
            db.import_verified([bad], source=SOURCE, file_name="x.txt", file_sha256="0" * 64, file_format="txt",
                               duplicates=0)
    with pytest.raises(InvalidRecordError, match="same number twice"):
        db.import_verified([good, good], source=SOURCE, file_name="x.txt", file_sha256="0" * 64,
                           file_format="txt", duplicates=0)
    with pytest.raises(InvalidRecordError, match="source"):
        db.import_verified([good], source=" ", file_name="x.txt", file_sha256="0" * 64, file_format="txt",
                           duplicates=0)
    assert db.imported_count() == 0


# --- PART C: existing data, provenance, page progress ------------------------------

def test_existing_database_is_preserved(db_path: Path, tmp_path: Path) -> None:
    with Database(db_path) as db:
        seed_website_page(db)
        before_records, before_pages, run_range = website_rows(db), page_rows(db), db.load_run_range()
    source_file = write(tmp_path, "list.txt", "1\n150\n250\n")
    with Database(db_path) as db:
        _, outcome = import_file(db, source_file, SOURCE)
        assert (outcome.inserted, outcome.already_scraped, outcome.already_imported) == (2, 1, 0)
        assert website_rows(db) == before_records and db.record_count() == 2
        assert page_rows(db) == before_pages and db.load_run_range() == run_range
        assert db.imported_row(1) is None and db.provenance(1) == WEBSITE

        first_import = db.imported_row(150)
        _, again = import_file(db, source_file, "second delivery")
        assert (again.inserted, again.already_scraped, again.already_imported) == (0, 1, 2)
        assert db.imported_row(150) == first_import  # never overwritten by a later import
        assert len(db.import_batches()) == 2


def test_import_never_marks_website_pages_completed(db: Database) -> None:
    seed_website_page(db)
    before = (page_rows(db), db.progress_counts(), db.pages_to_process(1, 3))
    import_file(db, FIXTURES / "verified_numbers.txt", SOURCE)  # numbers on pages 2, 3 and 671,089
    assert (page_rows(db), db.progress_counts(), db.pages_to_process(1, 3)) == before
    assert db.page(2)["status"] == "pending" and db.page(671_089) is None


def test_import_provenance_is_recorded_separately(db: Database) -> None:
    seed_website_page(db)
    verified, _ = import_file(db, FIXTURES / "verified_numbers.csv", SOURCE)
    batch = db.import_batches()[0]
    assert batch["source"] == SOURCE and batch["file_name"] == "verified_numbers.csv"
    assert batch["file_sha256"] == hashlib.sha256((FIXTURES / "verified_numbers.csv").read_bytes()).hexdigest()
    assert (batch["file_format"], batch["numbers_in_file"], batch["inserted"]) == ("csv", 2, 2)
    row = db.imported_row(150)
    assert row["import_id"] == batch["id"] and row["imported_at"] == batch["imported_at"]
    assert (db.provenance(2), db.provenance(150), db.provenance(151)) == (WEBSITE, IMPORTED, None)
    assert db.record_row(150) is None and [r.number for r in db.iter_records()] == [1, 2]
    assert verified.sha256 == batch["file_sha256"]


def test_later_website_scrape_keeps_both_provenances(db: Database) -> None:
    seed_website_page(db)
    import_file(db, FIXTURES / "verified_numbers.csv", SOURCE)
    run_pages(db, 1, 3, FakeSite(verified={2: [150]}), max_pages=1)
    assert db.page(2)["status"] == "completed" and db.provenance(150) == WEBSITE
    assert db.imported_row(150) is not None
    assert [r.number for r in db.iter_all_records()] == [1, 2, 150, 250]


def test_other_commands_do_not_create_import_tables(db: Database) -> None:
    seed_website_page(db)
    assert not db.has_imports() and db.imported_count() == 0
    assert [r.number for r in db.iter_all_records()] == [1, 2]


def test_import_cannot_be_combined_with_scraping(db_path: Path) -> None:
    for extra in (["--live", "--resume"], ["--start-page", "1", "--end-page", "2"], ["--resume"]):
        with pytest.raises(SystemExit) as exc:
            cli.main(["--db", str(db_path), "--import-verified", str(FIXTURES / "verified_numbers.txt"),
                      "--source", SOURCE, *extra])
        assert exc.value.code == 2
    with pytest.raises(SystemExit):
        cli.main(["--db", str(db_path), "--source", SOURCE])
    assert not db_path.exists()


def test_importer_has_no_network_code() -> None:
    source = (PROJECT_ROOT / "scraper" / "importer.py").read_text(encoding="utf-8")
    for word in ("httpx", "playwright", "urllib", "socket", "requests", "browser", "http://", "https://"):
        assert word not in source


# --- PART D: export ------------------------------------------------------------------

def test_csv_export_includes_imported_records(db: Database, tmp_path: Path) -> None:
    seed_website_page(db)
    import_file(db, FIXTURES / "verified_numbers.txt", SOURCE)
    out = tmp_path / "all.csv"
    assert export_from_db(db, out) == 5
    with out.open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    assert rows[0] == ["Number", "Start hex", "End hex", "Status"]
    assert [r[0] for r in rows[1:]] == ["1", "2", "150", "250", str(SEGMENT_COUNT)]
    assert rows[3] == ["150", "400009500000000000", "400009600000000000", "Проверено"]
    assert rows[-1] == [str(SEGMENT_COUNT), "7ffffff00000000000", "800000000000000000", "Проверено"]
    assert export_from_db(db, tmp_path / "web.csv", "website") == 2
    assert export_from_db(db, tmp_path / "imp.csv", "imported") == 3


def test_xlsx_export_keeps_hex_as_text(db: Database, tmp_path: Path) -> None:
    import_file(db, FIXTURES / "verified_numbers.txt", SOURCE)
    out = tmp_path / "all.xlsx"
    assert export_from_db(db, out) == 3
    ws = load_workbook(out)["Verified"]
    assert [c.value for c in ws[1]] == ["Number", "Start hex", "End hex", "Status"]
    assert [c.value for c in ws[2]] == [150, "400009500000000000", "400009600000000000", "Проверено"]
    for row in ws.iter_rows(min_row=2):
        for cell in row[1:3]:
            assert cell.data_type == "s" and cell.number_format == "@" and len(cell.value) == 18


def test_exported_csv_can_be_reimported_without_duplicates(db: Database, tmp_path: Path) -> None:
    import_file(db, FIXTURES / "verified_numbers.txt", SOURCE)
    out = tmp_path / "verified_records.csv"
    export_from_db(db, out)
    _, outcome = import_file(db, out, SOURCE)
    assert (outcome.inserted, outcome.already_imported) == (0, 3)


def test_cli_import_then_export(db_path: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    out = tmp_path / "imported.xlsx"
    assert cli.main(["--db", str(db_path), "--import-verified", str(FIXTURES / "numbers_only.csv"),
                     "--source", SOURCE, "--export", "xlsx", "--export-source", "imported",
                     "--output", str(out)]) == 0
    assert f"Exported 2 verified records to {out}" in capsys.readouterr().out
    assert cli.main(["--db", str(db_path), "--export", "csv", "--output", str(tmp_path / "all.csv")]) == 0
    assert "Exported 2 verified records" in capsys.readouterr().out
    with pytest.raises(SystemExit):
        cli.main(["--db", str(db_path), "--export-source", "imported"])
