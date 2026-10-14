"""Offline demonstration of the stored results; never contacts the website.

The database is opened read-only. The progress display at the end is simulated on a
temporary copy, so the real database is never written.
"""

from __future__ import annotations

import argparse
import csv
import re
import sqlite3
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from scraper.config import DATA_DIR, DB_PATH, HEX_WIDTH
from scraper.database import Database
from scraper.exporter import CSV_COLUMNS, TEXT_FORMAT
from scraper.parser import VERIFIED_TEXT, segment_range
from scraper.progress import ProgressReporter
from scraper.runner import plan_run

HEX_RE = re.compile(rf"[0-9a-f]{{{HEX_WIDTH}}}")
SIMULATED_RANGE = (350_000, 350_100)
SIMULATED_DURATIONS_S = (24.2, 17.5, 24.4)

COMMANDS = """\
  python main.py --start-page 350000 --end-page 350100          # planeja o intervalo, nada é baixado
  python main.py --start-page 350000 --end-page 350100 --live   # baixa 1 página do intervalo
  python main.py --resume --live --max-pages 20 --nav-timeout 60 # continua de onde parou (até 20 páginas)
  python main.py --export csv                                   # data/verified_records.csv
  python main.py --export xlsx                                  # data/verified_records.xlsx"""

Row = tuple[int, str, str, str]


def heading(title: str) -> None:
    print(f"\n=== {title} ===")


def check(ok: bool, label: str, failures: list[str]) -> None:
    print(f"  [{'OK' if ok else 'FALHA'}] {label}")
    if not ok:
        failures.append(label)


def read_csv(path: Path) -> tuple[tuple[str, ...], list[Row]]:
    with path.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        header = tuple(next(reader))
        return header, [(int(n), s, e, st) for n, s, e, st in reader]


def read_xlsx(path: Path) -> tuple[tuple[str, ...], list[Row], bool]:
    ws = load_workbook(path).active
    rows = list(ws.iter_rows())
    header = tuple(c.value for c in rows[0])
    data = [(int(r[0].value), r[1].value, r[2].value, r[3].value) for r in rows[1:]]
    hex_as_text = all(c.data_type == "s" and c.number_format == TEXT_FORMAT for r in rows[1:] for c in r[1:3])
    return header, data, hex_as_text


def show_database(db: Database, failures: list[str]) -> list[Row]:
    heading("Banco de dados SQLite (somente leitura)")
    check(db.conn.execute("PRAGMA integrity_check").fetchone()[0] == "ok", "integridade do SQLite", failures)
    run_range = db.load_run_range()
    print(f"  intervalo salvo: {run_range[0]}..{run_range[1]}" if run_range else "  intervalo salvo: nenhum")

    heading("Status das páginas")
    print(f"  {'página':>7}  {'status':<10} {'tentativas':>10} {'registros':>9}  atualizado em (UTC)")
    for p in db.conn.execute("SELECT * FROM page_progress ORDER BY page_number"):
        print(f"  {p['page_number']:>7}  {p['status']:<10} {p['attempts']:>10} "
              f"{p['records_found'] if p['records_found'] is not None else '-':>9}  {p['updated_at']}")

    heading("Registros verificados armazenados")
    rows = [(r.number, r.start_hex, r.end_hex, r.status) for r in db.iter_records()]
    print(f"  {'Number':>6}  {'Start hex':<18}  {'End hex':<18}  Status")
    for number, start, end, status in rows:
        print(f"  {number:>6}  {start}  {end}  {status}")
    print(f"  total: {len(rows)}")
    return rows


def check_hex(rows: list[Row], failures: list[str]) -> None:
    complete = all(HEX_RE.fullmatch(s) and HEX_RE.fullmatch(e) for _, s, e, _ in rows)
    check(complete, f"Start/End hex completos ({HEX_WIDTH} dígitos hexadecimais, sem máscara)", failures)
    check(all(segment_range(n) == (s, e) for n, s, e, _ in rows), "hex confere com o número do segmento", failures)
    check(all(st == VERIFIED_TEXT for *_, st in rows), f"todos com status '{VERIFIED_TEXT}'", failures)


def show_exports(csv_path: Path, xlsx_path: Path, reference: list[Row] | None, failures: list[str]) -> None:
    heading("Exportações CSV e Excel")
    csv_header, csv_rows = read_csv(csv_path)
    xlsx_header, xlsx_rows, hex_as_text = read_xlsx(xlsx_path)
    print(f"  {csv_path.name}: {len(csv_rows)} registros | {xlsx_path.name}: {len(xlsx_rows)} registros")
    check(csv_header == CSV_COLUMNS and xlsx_header == CSV_COLUMNS, f"colunas {', '.join(CSV_COLUMNS)}", failures)
    if reference is not None:
        check(csv_rows == reference, "CSV idêntico ao SQLite", failures)
        check(xlsx_rows == reference, "Excel idêntico ao SQLite", failures)
    else:
        check(csv_rows == xlsx_rows, "CSV idêntico ao Excel", failures)
    check(hex_as_text, "Start/End hex gravados como texto no Excel (formato '@')", failures)
    check_hex(csv_rows, failures)


def simulated_progress(db: Database | None) -> str:
    """Render the live progress display for a simulated run on a temporary database copy."""
    start, end = SIMULATED_RANGE
    with tempfile.TemporaryDirectory() as tmp:
        copy_path = Path(tmp) / "demo.db"
        if db is not None:
            target = sqlite3.connect(copy_path)
            db.conn.backup(target)
            target.close()
        now = [0.0]
        shown: list[str] = []
        with Database(copy_path) as copy:
            plan_run(copy, start, end)
            reporter = ProgressReporter(copy, start, end, page_limit=20, clock=lambda: now[0],
                                        notify=shown.append)
            page = start
            for page, seconds in zip(range(start, end + 1), SIMULATED_DURATIONS_S, strict=False):
                reporter.page_started(page)
                copy.claim_page(page)
                now[0] += seconds
                copy.complete_page(page, [], rows_inspected=100)
                reporter.page_finished(page, True)
            reporter.page_started(page + 1)
    return shown[-1]


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Demonstração offline dos resultados armazenados.")
    parser.add_argument("--db", type=Path, default=DB_PATH, help="banco SQLite (aberto somente para leitura)")
    parser.add_argument("--csv", type=Path, default=DATA_DIR / "verified_records.csv")
    parser.add_argument("--xlsx", type=Path, default=DATA_DIR / "verified_records.xlsx")
    args = parser.parse_args(argv)

    print("Demonstração do HonestPool Scraper (offline: nenhum acesso ao site)")
    failures: list[str] = []
    db = Database(args.db, readonly=True) if args.db.exists() else None
    try:
        reference = None
        if db is None:
            print(f"\n(banco {args.db} não incluído; a demonstração usa apenas as exportações de exemplo)")
        else:
            reference = show_database(db, failures)
            check_hex(reference, failures)
        show_exports(args.csv, args.xlsx, reference, failures)

        heading("Exemplos de comandos")
        print(COMMANDS)

        heading("Exemplo de exibição de progresso (simulação em cópia temporária, sem acesso ao site)")
        print(f"  tempos por página de {', '.join(f'{s} s' for s in SIMULATED_DURATIONS_S)}, "
              "baseados no tempo medido da página 2 e no intervalo configurado entre páginas")
        print(simulated_progress(db))
    finally:
        if db is not None:
            db.close()

    heading("Resultado")
    print("  todas as verificações passaram" if not failures else f"  {len(failures)} verificação(ões) falharam")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
