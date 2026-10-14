"""CSV and Excel export of verified records."""

from __future__ import annotations

import csv
import os
import tempfile
import time
from collections.abc import Iterable, Iterator
from contextlib import contextmanager
from pathlib import Path
from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from scraper.parser import SegmentRecord

if TYPE_CHECKING:
    from scraper.database import Database

CSV_COLUMNS = ("Number", "Start hex", "End hex", "Status")
XLSX_COLUMN_WIDTHS = {"A": 12, "B": 24, "C": 24, "D": 16}
XLSX_TEXT_COLUMNS = ("B", "C")
TEXT_FORMAT = "@"
SUPPORTED_FORMATS = ("csv", "xlsx")
EXPORT_SOURCES = ("all", "website", "imported")


def export_csv(records: Iterable[SegmentRecord], path: Path) -> int:
    """Write records sorted by Number; UTF-8 with BOM so Excel shows Cyrillic correctly."""
    rows = sorted(records, key=lambda r: r.number)
    with _atomic_output(path, "csv") as tmp:
        with tmp.open("w", encoding="utf-8-sig", newline="") as fh:
            writer = csv.writer(fh)
            writer.writerow(CSV_COLUMNS)
            for r in rows:
                writer.writerow((r.number, r.start_hex, r.end_hex, r.status))
            fh.flush()
            os.fsync(fh.fileno())  # on disk before the rename, so a power loss cannot leave a truncated CSV
    return len(rows)


def export_xlsx(records: Iterable[SegmentRecord], path: Path) -> int:
    """Write records to an .xlsx sheet with Start/End hex stored as text cells."""
    rows = sorted(records, key=lambda r: r.number)
    wb = Workbook()
    ws = wb.active
    ws.title = "Verified"
    ws.append(CSV_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="198754")
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append((r.number, r.start_hex, r.end_hex, r.status))
    for column in XLSX_TEXT_COLUMNS:
        for cell in ws[column]:
            # Typed as a string and formatted as Text, so Excel never reads it as a number.
            cell.value = str(cell.value)
            cell.data_type = "s"
            cell.number_format = TEXT_FORMAT
    for column, width in XLSX_COLUMN_WIDTHS.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    with _atomic_output(path, "xlsx") as tmp:
        wb.save(tmp)
    return len(rows)


def export_records(records: Iterable[SegmentRecord], path: Path) -> int:
    """Export by file suffix (.csv or .xlsx)."""
    fmt = path.suffix.lower().lstrip(".")
    if fmt == "csv":
        return export_csv(records, path)
    if fmt == "xlsx":
        return export_xlsx(records, path)
    raise ValueError(f"unsupported export format {path.suffix!r}; use one of {SUPPORTED_FORMATS}")


def export_from_db(db: Database, path: Path, source: str = "all") -> int:
    """Export website and imported records ("all"), or only one of them; the columns are the same."""
    readers = {"all": db.iter_all_records, "website": db.iter_records, "imported": db.iter_imported_records}
    if source not in readers:
        raise ValueError(f"unknown export source {source!r}; use one of {EXPORT_SOURCES}")
    return export_records(readers[source](), path)


def remove_stale_temporaries(path: Path, older_than_s: float = 60.0) -> int:
    """Delete temporary files that an interrupted export of ``path`` left behind; return how many."""
    path = Path(path)
    removed = 0
    cutoff = time.time() - older_than_s
    for tmp in path.parent.glob(f".{path.stem}.*{path.suffix}"):
        try:
            if tmp.is_file() and tmp.stat().st_mtime < cutoff:
                tmp.unlink()
                removed += 1
        except OSError:
            continue
    return removed


@contextmanager
def _atomic_output(path: Path, suffix: str) -> Iterator[Path]:
    """Yield a temporary path in the target directory and rename it over ``path`` on success."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix=f".{path.stem}.", suffix=f".{suffix}", dir=path.parent)
    os.close(fd)
    tmp = Path(name)
    try:
        yield tmp
        os.replace(tmp, path)
    finally:
        tmp.unlink(missing_ok=True)
